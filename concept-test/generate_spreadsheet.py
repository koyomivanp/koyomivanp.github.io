#!/usr/bin/env python3
"""Generate concept test survey data as XLSX (upload to Google Drive to convert to Sheets)."""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

HEADERS = [
    "タイムスタンプ",
    "学生区分",
    "性別",
    "現在使っている筆箱",
    "質問1-1",
    "質問1-2",
    "質問1-3",
    "質問1-4",
    "質問1-5",
]

ROWS = [
    ["", "中学生", "男性", "横型", 2, 3, 2, 2, 3],
    ["", "小学4〜6年生", "男性", "横型, ロールペンケース", 1, 5, 5, 1, 1],
    ["", "小学4〜6年生", "男性", "横型", 5, 5, 5, 5, 5],
    ["", "小学4〜6年生", "男性", "縦型", 2, 1, 1, 3, 2],
    ["", "小学1〜3年生", "女性", "横型, その他", 5, 5, 5, 5, 5],
    ["", "小学1〜3年生", "女性", "パカパカ筆箱, その他", 5, 5, 5, 5, 5],
    ["", "中学生", "回答しない", "縦型", 4, 4, 1, 4, 3],
    ["", "中学生", "", "横型, その他", 2, 1, 1, 1, 1],
    ["", "中学生", "女性", "横型", 2, 3, 5, 1, 1],
    ["", "中学生", "女性", "その他（ぬいぐるみ）", 5, 1, 1, 5, 1],
    ["", "中学生", "男性", "ロールペンケース", 4, 1, 1, 5, 3],
    ["", "中学生", "男性", "横型", 5, 5, 4, 3, 4],
]

OUTPUT = "26年度_商品開発_コンセプトテスト_紙回答.xlsx"


def main() -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "回答"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4285F4", end_color="4285F4", fill_type="solid")
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append(HEADERS)
    for row in ROWS:
        ws.append(row)

    for col in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for row in range(2, len(ROWS) + 2):
        for col in range(1, len(HEADERS) + 1):
            ws.cell(row=row, column=col).border = border
            if col >= 5:
                ws.cell(row=row, column=col).alignment = Alignment(horizontal="center")

    widths = [18, 14, 10, 24, 10, 10, 10, 10, 10]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = width

    ws.freeze_panes = "A2"
    wb.save(OUTPUT)
    print(f"Created: {OUTPUT}")


if __name__ == "__main__":
    main()
